"""Tests for evidence capture and attendance export. No camera or GPU needed.

Covers the two predecessor defects this phase exists to fix: an export that wrote
twenty rows of `Pending` with empty timestamps and told nobody, and an evidence
trail that consisted of the string "Hand only".

Run:  .venv\\Scripts\\python.exe tools\\test_export.py
"""

from __future__ import annotations

import csv
import json
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import numpy as np  # noqa: E402

from config import Config  # noqa: E402
from evidence import EvidenceStore  # noqa: E402
from export import CSV_HEADER, Exporter, ExportError, resolve_output_dir  # noqa: E402
from roster import ABSENT, PENDING, PRESENT, Student, assign_short_unique_calls  # noqa: E402

BASE = Config.load(ROOT / "config.json")

_results: list[tuple[bool | None, str, str]] = []


def check(cond: bool, name: str, detail: str = "") -> None:
    _results.append((bool(cond), name, detail))


def section(title: str) -> None:
    _results.append((None, title, ""))


# --- fixtures -------------------------------------------------------------
class FakeTrack:
    def __init__(self, tid=3, bbox=(200.0, 80.0, 320.0, 420.0), peak=0.87):
        self.track_id = tid
        self.bbox = bbox
        self.best_bbox = bbox
        self.peak_ema = peak


class FakeSnap:
    def __init__(self, frame):
        self.frame = frame


class FakeVideo:
    def __init__(self, frame=None):
        self._frame = frame

    def snapshot(self):
        return FakeSnap(self._frame)


def make_students():
    st = [Student(roll=f"SP25-BCS-{i + 1:03d}", name=f"Student {i + 1}") for i in range(5)]
    assign_short_unique_calls(st)
    st[0].status = PRESENT
    st[0].marked_at = "2026-09-27T10:00:01"
    st[0].evidence = "Hand raised"
    st[0].track_id = 3
    st[0].score = 0.871
    st[1].status = ABSENT
    st[1].marked_at = "2026-09-27T10:00:08"
    st[1].evidence = "Timeout: no hand raised"
    st[2].status = PRESENT
    st[2].marked_at = "2026-09-27T10:00:14"
    st[2].evidence = "Manual (teacher)"
    st[2].track_id = 7
    st[2].score = 0.64
    # st[3], st[4] deliberately left Pending
    return st


def cfg_with_dir(tmp: Path):
    return BASE.override({"output": {"dir": str(tmp)}})


# =========================================================================
def test_csv_contents() -> None:
    section("CSV export")
    tmp = Path(tempfile.mkdtemp())
    try:
        cfg = cfg_with_dir(tmp)
        ex = Exporter(cfg, session_id="TEST")
        students = make_students()
        paths = ex.export(students, telemetry=[{"roll": "SP25-BCS-001", "track_id": 3}])
        csv_path = next(p for p in paths if p.suffix == ".csv")
        check(csv_path.exists(), "csv written", csv_path.name)

        raw = csv_path.read_bytes()
        check(raw.startswith(b"\xef\xbb\xbf"), "written as utf-8-sig so Excel reads it")

        with open(csv_path, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh))
        check(rows[0] == CSV_HEADER, "header matches the documented schema", str(rows[0][:5]))
        check(
            rows[0][:5] == ["Roll Number", "Name", "Status", "Marked At", "Evidence"],
            "first five columns match the predecessor's for compatibility",
        )
        check(len(rows) == 6, "one row per student plus header", f"{len(rows)} rows")
        check(rows[1][2] == PRESENT and rows[1][5] == "3", "present row carries its track id",
              str(rows[1][2:6]))
        check(rows[4][2] == PENDING, "pending students are exported AS Pending, not Absent",
              rows[4][2])
        check(rows[4][3] == "", "pending rows have no fabricated timestamp")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_pending_guard() -> None:
    section("pending guard")
    tmp = Path(tempfile.mkdtemp())
    try:
        ex = Exporter(cfg_with_dir(tmp), session_id="TEST")
        students = make_students()
        check(ex.pending_count(students) == 2, "pending count reported to the caller",
              str(ex.pending_count(students)))
        # The predecessor silently wrote an all-Pending file. Here the count is
        # surfaced so the UI can ask, and it is recorded in the session file.
        paths = ex.export(students)
        sess = json.loads(next(p for p in paths if p.name.startswith("session_")).read_text(encoding="utf-8"))
        check(sess["summary"]["pending"] == 2, "pending count recorded in the session file",
              str(sess["summary"]))
        check(
            any("never called" in str(v) for v in sess.get("summary", {}).values())
            or sess["summary"]["pending"] == 2,
            "unfinished roster is visible in the export",
        )
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_empty_roster_refused() -> None:
    section("empty roster")
    tmp = Path(tempfile.mkdtemp())
    try:
        ex = Exporter(cfg_with_dir(tmp), session_id="TEST")
        try:
            ex.export([])
            check(False, "exporting nothing raises")
        except ExportError:
            check(True, "exporting nothing raises ExportError")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_xlsx() -> None:
    section("XLSX export")
    tmp = Path(tempfile.mkdtemp())
    try:
        ex = Exporter(cfg_with_dir(tmp), session_id="TEST")
        students = make_students()
        paths = ex.export(students)
        xlsx = [p for p in paths if p.suffix == ".xlsx"]
        if not xlsx:
            check(False, "xlsx written (openpyxl missing?)")
            return
        check(xlsx[0].exists(), "xlsx written", xlsx[0].name)

        from openpyxl import load_workbook

        wb = load_workbook(xlsx[0])
        check("Attendance" in wb.sheetnames and "Summary" in wb.sheetnames,
              "has Attendance and Summary sheets", str(wb.sheetnames))
        ws = wb["Attendance"]
        check(ws.max_row == 6, "one row per student plus header", f"{ws.max_row}")
        check(ws.freeze_panes == "A2", "header row frozen")
        summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
        check(summary.get("Present") == 2, "summary counts present", str(summary.get("Present")))
        check(summary.get("Pending") == 2, "summary counts pending", str(summary.get("Pending")))
        check("WARNING" in summary, "summary warns about students never called")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_session_json_and_track_audit() -> None:
    section("session telemetry")
    tmp = Path(tempfile.mkdtemp())
    try:
        ex = Exporter(cfg_with_dir(tmp), session_id="TEST")
        students = make_students()
        telemetry = [
            {"roll": "SP25-BCS-001", "status": PRESENT, "track_id": 3,
             "contenders": [3], "confirmed": [3], "seconds_to_decide": 1.2},
            {"roll": "SP25-BCS-002", "status": ABSENT, "track_id": None,
             "contenders": [], "confirmed": [], "seconds_to_decide": 6.4},
        ]
        paths = ex.export(students, telemetry=telemetry,
                          runtime={"model": "yolo11m-pose.pt", "fps": 15.2})
        sp = next(p for p in paths if p.name.startswith("session_"))
        data = json.loads(sp.read_text(encoding="utf-8"))
        check(len(data["marks"]) == 2, "per-mark telemetry recorded", str(len(data["marks"])))
        check(data["runtime"]["model"] == "yolo11m-pose.pt", "runtime details recorded")
        check("temporal" in data["config"], "full config snapshot recorded for later tuning")
        check(data["marks"][0]["seconds_to_decide"] == 1.2, "decision latency recorded")
        check(data["track_reuse"] == {}, "no track reuse in a clean run", str(data["track_reuse"]))

        # The audit that would have answered the live-session question.
        students[3].status = PRESENT
        students[3].track_id = 3  # same tracked person credited for a second roll
        paths = ex.export(students)
        sp = next(p for p in paths if p.name.startswith("session_"))
        data = json.loads(sp.read_text(encoding="utf-8"))
        check("3" in data["track_reuse"], "track reuse IS detected and reported",
              str(data["track_reuse"]))
        check(len(data["track_reuse"]["3"]) == 2, "names both rolls involved",
              str(data["track_reuse"].get("3")))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_evidence_capture() -> None:
    section("evidence images")
    tmp = Path(tempfile.mkdtemp())
    try:
        cfg = cfg_with_dir(tmp)
        store = EvidenceStore(cfg, session_id="TEST")
        frame = np.random.randint(40, 200, (480, 640, 3), dtype=np.uint8)
        video = FakeVideo(frame)
        s = Student(roll="SP25-BCS-007", name="Ayesha")

        path = store.capture(s, FakeTrack(), video)
        check(bool(path), "evidence image written", path)
        check(not Path(path).is_absolute(),
              "path is relative so the folder stays portable", path)
        p = store.base / path
        check(p.exists() and p.stat().st_size > 1000, "file is a real image",
              f"{p.stat().st_size if p.exists() else 0} bytes")
        check("SP25-BCS-007" in p.name, "named after the roll number", p.name)

        import cv2

        img = cv2.imread(str(p))
        check(img is not None, "image decodes")
        # Crop is 120x340 plus padding, plus a caption strip.
        check(img is not None and img.shape[0] > 340, "caption strip added below the crop",
              str(None if img is None else img.shape))

        # A tiny back-row crop is upscaled to stay viewable.
        small = store.capture(
            Student(roll="SP25-BCS-008"), FakeTrack(tid=9, bbox=(10.0, 10.0, 60.0, 120.0)), video
        )
        img2 = cv2.imread(str(store.base / small))
        check(img2 is not None and img2.shape[1] >= 240, "small crops upscaled for legibility",
              str(None if img2 is None else img2.shape))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_evidence_never_raises() -> None:
    section("evidence failure modes")
    tmp = Path(tempfile.mkdtemp())
    try:
        store = EvidenceStore(cfg_with_dir(tmp), session_id="TEST")
        s = Student(roll="SP25-BCS-009")
        # A mark must never fail because evidence could not be written.
        check(store.capture(s, None, FakeVideo(None)) == "", "no track -> empty, no raise")
        check(store.capture(s, FakeTrack(), FakeVideo(None)) == "", "no frame -> empty, no raise")
        check(store.capture(s, FakeTrack(), None) == "", "no video -> empty, no raise")

        class Boom:
            def snapshot(self):
                raise RuntimeError("camera exploded")

        check(store.capture(s, FakeTrack(), Boom()) == "", "exploding video -> empty, no raise")

        off = EvidenceStore(
            BASE.override({"output": {"dir": str(tmp), "save_evidence_images": False}}),
            session_id="TEST",
        )
        frame = np.zeros((480, 640, 3), dtype=np.uint8)
        check(off.capture(s, FakeTrack(), FakeVideo(frame)) == "",
              "disabled by config -> nothing written")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_output_dir_default() -> None:
    section("output location")
    d = resolve_output_dir(BASE)
    check(d.name == "SmartRollCall Attendance",
          "defaults to the predecessor's folder so existing habits keep working", str(d))
    custom = resolve_output_dir(BASE.override({"output": {"dir": r"C:\tmp\x"}}))
    check(str(custom).endswith("x"), "honours a configured directory", str(custom))


def main() -> int:
    for fn in (
        test_csv_contents,
        test_pending_guard,
        test_empty_roster_refused,
        test_xlsx,
        test_session_json_and_track_audit,
        test_evidence_capture,
        test_evidence_never_raises,
        test_output_dir_default,
    ):
        try:
            fn()
        except Exception as exc:
            check(False, f"{fn.__name__} crashed", f"{type(exc).__name__}: {exc}")

    print()
    print("=" * 74)
    passed = failed = 0
    for ok, name, detail in _results:
        if ok is None:
            print(f"\n  -- {name} --")
            continue
        suffix = f"   [{detail}]" if detail else ""
        print(f"  {'PASS' if ok else 'FAIL'}  {name}{suffix}")
        if ok:
            passed += 1
        else:
            failed += 1
    print("=" * 74)
    print(f"  {passed} passed, {failed} failed")
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
